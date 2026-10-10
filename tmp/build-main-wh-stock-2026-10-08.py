"""Main-warehouse stock Excel as of 2026-10-08 15:55 Asia/Colombo."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.legend import Legend
from openpyxl.chart.label import DataLabelList
from openpyxl.chart import BarChart, Reference

AGENT = Path(r"C:\Users\Bad-Boy\.cursor\projects\c-dev-cosmo-os\agent-tools")
OUT = Path(r"c:\dev\cosmo-os\tmp\main-warehouse-stock-2026-10-08-1555.xlsx")
SLE_PATH = Path(r"c:\dev\cosmo-os\tmp\sle-after-1555-2026-10-08.json")

SITES = [
    {
        "sheet": "Cosmetics",
        "label": "Cosmetics",
        "company": "Cosmetics.lk",
        "warehouse": "Main Warehouse - Cosmo",
        "erp": "cosmetics-lk-01",
        "file": "802bb8dc-38f8-492f-a8b0-0b2645684aa0.txt",
    },
    {
        "sheet": "SPK",
        "label": "SPK",
        "company": "SPK Trading (Pvt) Ltd",
        "warehouse": "Main Warehouse - SPK",
        "erp": "cosmetics-lk-02",
        "file": "76bfa2e9-784b-4d60-926f-40015469494b.txt",
    },
    {
        "sheet": "PEVI",
        "label": "PEVI",
        "company": "Pevi Trading Lanka (Pvt) Ltd",
        "warehouse": "Main Warehouses - Pevi",
        "erp": "cosmetics-lk-02",
        "file": "c24c29ae-0aac-415c-ad82-b7f3a0387ac7.txt",
    },
    {
        "sheet": "CAHMI",
        "label": "CAHMI",
        "company": "Chami Trading Lanka (Pvt) Ltd",
        "warehouse": "Main Warehouse - Chami",
        "erp": "cosmetics-lk-02",
        "file": "f5a3939c-4eeb-4857-a7a9-15c40883cf9f.txt",
    },
    {
        "sheet": "USP",
        "label": "USP",
        "company": "USP Trading (Pvt) Ltd",
        "warehouse": "Main Warehouse - USP",
        "erp": "cosmetics-lk-02",
        "file": "9a0fc643-ccd1-48a6-82cc-7acfeae814db.txt",
    },
    {
        "sheet": "UDARA",
        "label": "UDARA",
        "company": "Udara Trading Lanka (Pvt) Ltd",
        "warehouse": "Main Warehouse - Udara",
        "erp": "cosmetics-lk-02",
        "file": "66e5a47a-b9ca-4022-bdb1-035b5ab8a2b5.txt",
    },
    {
        "sheet": "KAD",
        "label": "KAD",
        "company": "KAD Trading (Pvt) Ltd",
        "warehouse": "Main Warehouse - KAD",
        "erp": "cosmetics-lk-02",
        "file": "111ce2a6-208f-46b0-991c-ca7b707113f0.txt",
    },
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F4E79")
SUB_FONT = Font(name="Arial", size=10, color="333333")
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="555555")
INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="000000")
TOTAL_FONT = Font(name="Arial", bold=True, size=10)
LINK_FONT = Font(name="Arial", size=10, color="008000")
THIN = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)
TOTAL_BORDER = Border(
    top=Side(style="medium", color="1F4E79"),
    bottom=Side(style="double", color="1F4E79"),
)
ZEBRA = PatternFill("solid", fgColor="F7F9FC")
NEG_FILL = PatternFill("solid", fgColor="FDECEC")
QTY_FMT = "#,##0.000;(#,##0.000);-"
VAL_FMT = '#,##0.00;(#,##0.00);"-"'
RATE_FMT = '#,##0.00;(#,##0.00);"-"'


def load_balance(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows = payload["message"]["result"]
    if not isinstance(rows, list):
        raise SystemExit(f"No result rows in {path}")
    return rows


def aggregate(rows: list[dict], warehouse: str) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        if row.get("warehouse") != warehouse:
            continue
        code = row["item_code"]
        slot = out.get(code)
        if slot is None:
            out[code] = {
                "item_code": code,
                "item_name": row.get("item_name") or "",
                "item_group": row.get("item_group") or "",
                "stock_uom": row.get("stock_uom") or "",
                "bal_qty": float(row.get("bal_qty") or 0),
                "bal_val": float(row.get("bal_val") or 0),
            }
        else:
            slot["bal_qty"] += float(row.get("bal_qty") or 0)
            slot["bal_val"] += float(row.get("bal_val") or 0)
    return out


def main() -> None:
    sle_rows = json.loads(SLE_PATH.read_text(encoding="utf-8"))
    sle_by_wh: dict[str, dict[str, list]] = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0]))
    for entry in sle_rows:
        bucket = sle_by_wh[entry["warehouse"]][entry["item_code"]]
        bucket[0] += float(entry["actual_qty"])
        bucket[1] += float(entry["stock_value_difference"])

    built: list[tuple[dict, list[dict]]] = []
    missing: list[str] = []
    negatives: list[str] = []

    for site in SITES:
        balance = aggregate(load_balance(AGENT / site["file"]), site["warehouse"])
        adjustments = sle_by_wh.get(site["warehouse"], {})
        for code, (qty_delta, val_delta) in adjustments.items():
            if code not in balance:
                missing.append(f"{site['sheet']} {code} after-qty {qty_delta}")
                balance[code] = {
                    "item_code": code,
                    "item_name": "",
                    "item_group": "",
                    "stock_uom": "",
                    "bal_qty": 0.0,
                    "bal_val": 0.0,
                }
        lines = []
        for code, row in balance.items():
            delta_qty, delta_val = adjustments.get(code, (0.0, 0.0))
            qty = round(row["bal_qty"] - delta_qty, 3)
            val = round(row["bal_val"] - delta_val, 2)
            if abs(qty) < 0.0005:
                continue
            if qty < 0:
                negatives.append(f"{site['sheet']} {code} qty {qty}")
            lines.append({**row, "qty": qty, "val": val, "adjusted": abs(delta_qty) > 0.0005})
        lines.sort(key=lambda r: r["item_code"])
        built.append((site, lines))
        print(f"{site['sheet']}: {len(lines)} items with qty, adjustments {len(adjustments)}")

    if missing:
        print("MISSING FROM BALANCE:")
        for line in missing:
            print(" ", line)
    if negatives:
        print("NEGATIVE:")
        for line in negatives:
            print(" ", line)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_summary_shell(summary)

    ranges: list[tuple[dict, int]] = []
    for site, lines in built:
        last = write_stock_sheet(wb.create_sheet(site["sheet"]), site, lines)
        ranges.append((site, last))

    write_adjustments(wb.create_sheet("After 1555"), sle_rows)
    fill_summary(summary, ranges)
    write_chart(summary, len(ranges))

    summary.sheet_view.showGridLines = False
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.oddHeader.left.text = "Main warehouse stock"
    summary.oddFooter.right.text = "Page &P of &N"
    summary.print_title_rows = "1:6"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


def write_summary_shell(ws) -> None:
    ws.sheet_properties.tabColor = "1F4E79"
    ws["A1"] = "Main warehouse stock"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = "As of 8 October 2026, 15:55 (Asia/Colombo)"
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A2:F2")
    ws["A3"] = (
        "Qty and stock value = ERPNext Stock Balance closing on 2026-10-08, "
        "minus Stock Ledger Entry qty and stock value difference posted after 15:55:00 that day. "
        "Rows with zero qty omitted. Valuation rate = stock value / qty."
    )
    ws["A3"].font = NOTE_FONT
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 36
    ws["A4"] = (
        "Date read as 8 October 2026 (10/08 month/day), the day before this file was pulled. "
        "Cosmetics from cosmetics-lk-01. SPK, PEVI, CAHMI, USP, UDARA, KAD from cosmetics-lk-02. "
        "PEVI warehouse name in ERP is Main Warehouses - Pevi. CAHMI warehouse is Main Warehouse - Chami."
    )
    ws["A4"].font = NOTE_FONT
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A4:F4")
    ws.row_dimensions[4].height = 36

    headers = ["Company", "Warehouse", "ERP site", "Items", "Total qty", "Stock value (LKR)"]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(6, col, text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
    ws.row_dimensions[6].height = 22
    ws.auto_filter.ref = "A6:F13"
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = "A6:F13"
    widths = [34, 28, 22, 12, 16, 22]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "Main warehouse stock 8 Oct 2026 15:55"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_title_rows = "1:6"
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110


def fill_summary(ws, ranges: list[tuple[dict, int]]) -> None:
    for i, (site, last) in enumerate(ranges):
        row = 7 + i
        ws.cell(row, 1, site["company"]).font = Font(name="Arial", size=10)
        ws.cell(row, 2, site["warehouse"]).font = Font(name="Arial", size=10)
        ws.cell(row, 3, site["erp"]).font = Font(name="Arial", size=10)
        item_cell = ws.cell(row, 4, f"=COUNTA({site['sheet']}!A5:A{last})")
        qty_cell = ws.cell(row, 5, f"=SUM({site['sheet']}!E5:E{last})")
        val_cell = ws.cell(row, 6, f"=SUM({site['sheet']}!F5:F{last})")
        for cell in (item_cell, qty_cell, val_cell):
            cell.font = LINK_FONT
        item_cell.number_format = "#,##0"
        qty_cell.number_format = QTY_FMT
        val_cell.number_format = VAL_FMT
        fill = ZEBRA if i % 2 else PatternFill()
        for col in range(1, 7):
            ws.cell(row, col).fill = fill
            ws.cell(row, col).border = THIN
            ws.cell(row, col).alignment = Alignment(vertical="center")
    total_row = 7 + len(ranges)
    ws.cell(total_row, 1, "Total").font = TOTAL_FONT
    ws.cell(total_row, 4, f"=SUM(D7:D{total_row - 1})").font = TOTAL_FONT
    ws.cell(total_row, 5, f"=SUM(E7:E{total_row - 1})").font = TOTAL_FONT
    ws.cell(total_row, 6, f"=SUM(F7:F{total_row - 1})").font = TOTAL_FONT
    ws.cell(total_row, 4).number_format = "#,##0"
    ws.cell(total_row, 5).number_format = QTY_FMT
    ws.cell(total_row, 6).number_format = VAL_FMT
    for col in range(1, 7):
        ws.cell(total_row, col).border = TOTAL_BORDER
        ws.cell(total_row, col).font = TOTAL_FONT
    note_row = total_row + 2
    ws.cell(
        note_row,
        1,
        "After 15:55 sheet lists every Stock Ledger Entry removed from the end-of-day balance. "
        "Blue numbers are ERP figures. Green numbers link to warehouse sheets. Black numbers are formulas.",
    ).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[note_row].height = 32
    # chart data is the company rows; keep filter covering data + header only
    ws.auto_filter.ref = f"A6:F{total_row - 1}"


def write_stock_sheet(ws, site: dict, lines: list[dict]) -> int:
    ws.sheet_properties.tabColor = "2E75B6"
    ws["A1"] = f"{site['label']} main warehouse stock"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")
    ws["A2"] = (
        f"{site['warehouse']}  |  {site['company']}  |  {site['erp']}  |  "
        "as of 8 October 2026 15:55 Asia/Colombo"
    )
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A2:G2")
    headers = [
        "Item code",
        "Item name",
        "Item group",
        "UOM",
        "Qty",
        "Stock value (LKR)",
        "Valuation rate (LKR)",
    ]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(4, col, text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:G{4 + max(len(lines), 1)}"

    for i, line in enumerate(lines):
        row = 5 + i
        values = [
            line["item_code"],
            line["item_name"],
            line["item_group"],
            line["stock_uom"],
            line["qty"],
            line["val"],
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.font = INPUT_FONT if col in (5, 6) else Font(name="Arial", size=10)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if i % 2:
                cell.fill = ZEBRA
        rate = ws.cell(row, 7, f'=IF(E{row}=0,0,F{row}/E{row})')
        rate.font = FORMULA_FONT
        rate.border = THIN
        rate.number_format = RATE_FMT
        rate.alignment = Alignment(vertical="center")
        ws.cell(row, 5).number_format = QTY_FMT
        ws.cell(row, 6).number_format = VAL_FMT
        if line["qty"] < 0:
            for col in range(1, 8):
                ws.cell(row, col).fill = NEG_FILL
        elif i % 2:
            rate.fill = ZEBRA

    last = 4 + len(lines)
    total_row = last + 1
    if lines:
        ws.cell(total_row, 1, "Total").font = TOTAL_FONT
        qty = ws.cell(total_row, 5, f"=SUM(E5:E{last})")
        val = ws.cell(total_row, 6, f"=SUM(F5:F{last})")
        rate = ws.cell(total_row, 7, f"=IF(E{total_row}=0,0,F{total_row}/E{total_row})")
        for cell, fmt in ((qty, QTY_FMT), (val, VAL_FMT), (rate, RATE_FMT)):
            cell.font = TOTAL_FONT
            cell.number_format = fmt
        for col in range(1, 8):
            ws.cell(total_row, col).font = TOTAL_FONT
            ws.cell(total_row, col).border = TOTAL_BORDER
    else:
        ws.cell(5, 1, "No stock").font = NOTE_FONT
        last = 5

    widths = [18, 52, 22, 10, 14, 20, 22]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = False
    ws.print_title_rows = "1:4"
    ws.page_setup.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.oddHeader.left.text = f"{site['label']} main stock 8 Oct 2026 15:55"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_view.zoomScale = 110
    ws.auto_filter.ref = f"A4:G{last}"
    return last


def write_adjustments(ws, sle_rows: list[dict]) -> None:
    ws.sheet_properties.tabColor = "C65911"
    ws["A1"] = "Ledger entries after 15:55 removed from end-of-day qty"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = "Stock Ledger Entry on 2026-10-08 with posting_time later than 15:55:00. Not cancelled."
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:F2")
    headers = ["Warehouse", "Item code", "Posting time", "Qty change", "Value change (LKR)", "SLE name"]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(4, col, text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = Alignment(horizontal="center")
    ordered = sorted(sle_rows, key=lambda r: (r["warehouse"], r["posting_time"], r["name"]))
    for i, entry in enumerate(ordered):
        row = 5 + i
        vals = [
            entry["warehouse"],
            entry["item_code"],
            entry["posting_time"],
            entry["actual_qty"],
            entry["stock_value_difference"],
            entry["name"],
        ]
        for col, value in enumerate(vals, 1):
            cell = ws.cell(row, col, value)
            cell.font = INPUT_FONT if col in (4, 5) else Font(name="Arial", size=10)
            cell.border = THIN
            if i % 2:
                cell.fill = ZEBRA
        ws.cell(row, 4).number_format = QTY_FMT
        ws.cell(row, 5).number_format = VAL_FMT
    last = 4 + len(ordered)
    total_row = last + 1
    ws.cell(total_row, 1, "Total").font = TOTAL_FONT
    q = ws.cell(total_row, 4, f"=SUM(D5:D{last})")
    v = ws.cell(total_row, 5, f"=SUM(E5:E{last})")
    q.number_format = QTY_FMT
    v.number_format = VAL_FMT
    for col in range(1, 7):
        ws.cell(total_row, col).font = TOTAL_FONT
        ws.cell(total_row, col).border = TOTAL_BORDER
    ws.auto_filter.ref = f"A4:F{last}"
    ws.freeze_panes = "A5"
    widths = [28, 16, 22, 14, 22, 24]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"
    ws.oddHeader.left.text = "SLE after 15:55 on 8 Oct 2026"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5


def write_chart(ws, n: int) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.title = "Stock value (LKR) by warehouse"
    chart.y_axis.title = "LKR"
    chart.x_axis.title = None
    data = Reference(ws, min_col=6, min_row=6, max_row=6 + n)
    cats = Reference(ws, min_col=1, min_row=7, max_row=6 + n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.style = 10
    chart.y_axis.numFmt = "#,##0"
    chart.width = 18
    chart.height = 8
    ws.add_chart(chart, "A20")


if __name__ == "__main__":
    main()
