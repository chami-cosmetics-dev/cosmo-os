"""Blank Purchases template for the one-time Vault supplement contact import."""

from datetime import date
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

out = r"c:\dev\cosmo-os\specs\062-vault-supplement-contacts\vault-supplement-contacts-template.xlsx"

headers = [
    ("customer_name", "yes", "Customer name", "Nimal Perera"),
    ("phone", "one", "Primary phone. Fill this or email", "0771234567"),
    ("email", "one", "Primary email. Fill this or phone", "nimal@example.com"),
    ("address", "no", "Street address", "12 Galle Rd"),
    ("city", "no", "City", "Colombo"),
    ("district", "no", "District", "Colombo"),
    ("invoice_no", "yes", "Invoice or order number shown on purchase history", "INV-1001"),
    ("invoice_date", "yes", "Purchase date as YYYY-MM-DD", "2024-03-15"),
    ("item_code", "yes", "Item code that exists on the Supplement Vault ERP item list", "SV-WHEY-1KG"),
    ("item_name", "yes", "Item name shown on the line", "Whey Protein 1kg"),
    ("quantity", "yes", "Quantity greater than zero", 2),
    ("unit_price", "yes", "Price per unit", 15000),
    ("line_amount", "no", "Line total. Blank means quantity times unit price", None),
    ("currency", "no", "Blank means LKR", "LKR"),
    ("payment_method", "no", "How the purchase was paid", "Cash"),
    ("location", "no", "Store or channel name", "Website"),
    ("merchant", "no", "Sales person name", "Kamal"),
    ("source_ref", "yes", "Stable Cosmo purchase id, or a unique id you assign for a hand-typed row", "COSMO-ORD-1001"),
]

header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
required_fill = PatternFill("solid", fgColor="1F4E79")
one_fill = PatternFill("solid", fgColor="C65911")
optional_fill = PatternFill("solid", fgColor="5B9BD5")
example_fill = PatternFill("solid", fgColor="FCE4D6")
input_font = Font(name="Arial", size=11)
title_font = Font(name="Arial", bold=True, size=16, color="1F4E79")
section_font = Font(name="Arial", bold=True, size=12, color="1F4E79")
body_font = Font(name="Arial", size=11)
thin = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
wrap = Alignment(wrap_text=True, vertical="center")

wb = Workbook()
purchases = wb.active
purchases.title = "Purchases"

header_fills = {"yes": required_fill, "one": one_fill, "no": optional_fill}
for col, (name, requirement, _meaning, _example) in enumerate(headers, start=1):
    cell = purchases.cell(1, col, name)
    cell.font = header_font
    cell.fill = header_fills[requirement]
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin

example = [
    "Nimal Perera",
    "0771234567",
    "nimal@example.com",
    "12 Galle Rd",
    "Colombo",
    "Colombo",
    "INV-1001",
    date(2024, 3, 15),
    "SV-WHEY-1KG",
    "Whey Protein 1kg",
    2,
    15000,
    "=K2*L2",
    "LKR",
    "Cash",
    "Website",
    "Kamal",
    "COSMO-ORD-1001",
]
for col, value in enumerate(example, start=1):
    cell = purchases.cell(2, col, value)
    cell.font = input_font
    cell.fill = example_fill
    cell.border = thin
    cell.alignment = Alignment(vertical="center")

purchases["B2"].number_format = "@"
purchases["H2"].number_format = "YYYY-MM-DD"
purchases["I2"].number_format = "@"
purchases["K2"].number_format = "#,##0.00"
purchases["L2"].number_format = "#,##0.00"
purchases["M2"].number_format = "#,##0.00"
purchases["A2"].comment = Comment(
    "EXAMPLE ROW. Delete this row before you upload the file. Orange cells are a sample, not a real customer.",
    "Cosmo OS",
    width=280,
    height=60,
)

widths = [22, 16, 28, 22, 16, 16, 16, 16, 18, 24, 12, 14, 14, 12, 18, 16, 16, 22]
for i, width in enumerate(widths, start=1):
    purchases.column_dimensions[get_column_letter(i)].width = width

purchases.row_dimensions[1].height = 30
purchases.freeze_panes = "A2"
purchases.auto_filter.ref = "A1:R2"
purchases.page_setup.orientation = "landscape"
purchases.page_setup.fitToPage = True
purchases.page_setup.fitToWidth = 1
purchases.page_setup.fitToHeight = 0
purchases.page_setup.paperSize = purchases.PAPERSIZE_A4
purchases.sheet_properties.tabColor = "1F4E79"
purchases.oddHeader.left.text = "Vault supplement contact import"
purchases.oddFooter.left.text = "Delete the orange example row before upload"
purchases.oddFooter.right.text = "Page &P of &N"
purchases.print_title_rows = "1:1"
purchases.page_setup.horizontalCentered = True
purchases.sheet_view.showGridLines = False
purchases.sheet_view.zoomScale = 110

# Keep phone and item code as text on the rows staff will type.
phone_text = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="40", allow_blank=True)
phone_text.add("B3:B5000")
purchases.add_data_validation(phone_text)
for row in range(3, 21):
    purchases.cell(row, 2).number_format = "@"
    purchases.cell(row, 8).number_format = "YYYY-MM-DD"
    purchases.cell(row, 9).number_format = "@"
    purchases.cell(row, 11).number_format = "#,##0.00"
    purchases.cell(row, 12).number_format = "#,##0.00"
    purchases.cell(row, 13).number_format = "#,##0.00"
    for col in range(1, 19):
        purchases.cell(row, col).font = input_font

guide = wb.create_sheet("Guide")
guide.sheet_properties.tabColor = "5B9BD5"
guide.sheet_view.showGridLines = False
guide["A1"] = "Vault supplement contact import"
guide["A1"].font = title_font
guide["A2"] = "One row on Purchases is one supplement item line for one customer. Fill Purchases, delete the orange example row, then upload the file to Vault OS."
guide["A2"].font = body_font
guide["A2"].alignment = wrap
guide.merge_cells("A2:D2")
guide.row_dimensions[2].height = 36

guide["A4"] = "Before you upload"
guide["A4"].font = section_font
rules = [
    "Delete the orange example row on Purchases. It is a sample customer, not someone to import.",
    "Dark blue headers are required on every row. Orange headers mean phone or email: fill at least one. Light blue headers are optional.",
    "A row with neither phone nor email is rejected.",
    "invoice_date must be a real date written YYYY-MM-DD.",
    "item_code must already exist on the Supplement Vault ERP item list. Cosmetics codes are rejected.",
    "quantity must be greater than zero.",
    "Leave line_amount blank to use quantity times unit price. The example row shows that formula.",
    "Leave currency blank to use LKR.",
    "source_ref plus invoice_no plus item_code identifies the line. Uploading the same three again updates that line.",
    "Same customer on several lines: repeat the customer columns on each item row.",
    "Do not rename headers and do not change their order.",
]
for i, rule in enumerate(rules, start=5):
    cell = guide.cell(i, 1, f"{i - 4}. {rule}")
    cell.font = body_font
    cell.alignment = wrap
    guide.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
    guide.row_dimensions[i].height = 32

start = 17
guide.cell(start, 1, "Header").font = header_font
guide.cell(start, 2, "Required").font = header_font
guide.cell(start, 3, "What to enter").font = header_font
guide.cell(start, 4, "Example").font = header_font
for col in range(1, 5):
    guide.cell(start, col).fill = required_fill
    guide.cell(start, col).alignment = Alignment(horizontal="center")
    guide.cell(start, col).border = thin

requirement_label = {"yes": "Yes", "one": "One of", "no": "No"}
for offset, (name, requirement, meaning, sample) in enumerate(headers, start=1):
    row = start + offset
    values = [name, requirement_label[requirement], meaning, "" if sample is None else sample]
    for col, value in enumerate(values, start=1):
        cell = guide.cell(row, col, value)
        cell.font = body_font
        cell.border = thin
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if name in ("phone", "item_code"):
            cell.number_format = "@"
    guide.cell(row, 2).fill = header_fills[requirement]
    guide.cell(row, 2).font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    guide.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center")
    if requirement == "no":
        guide.cell(row, 1).fill = PatternFill("solid", fgColor="D6EAF8")
    elif requirement == "one":
        guide.cell(row, 1).fill = PatternFill("solid", fgColor="FCE4D6")

guide.column_dimensions["A"].width = 24
guide.column_dimensions["B"].width = 14
guide.column_dimensions["C"].width = 78
guide.column_dimensions["D"].width = 28
guide.row_dimensions[1].height = 24
guide.page_setup.orientation = "landscape"
guide.page_setup.fitToPage = True
guide.page_setup.fitToWidth = 1
guide.page_setup.fitToHeight = 1
guide.page_setup.paperSize = guide.PAPERSIZE_A4
guide.page_setup.horizontalCentered = True
guide.oddHeader.left.text = "How to fill the template"
guide.oddFooter.right.text = "Page &P of &N"
guide.print_title_rows = "1:1"
guide.freeze_panes = "A18"
guide.auto_filter.ref = "A17:D35"
guide.page_setup.horizontalCentered = True

wb.properties.title = "Vault supplement contact import template"
wb.properties.subject = "One-time import of supplement customers and purchase lines into Vault OS"
wb.properties.creator = "Cosmo OS"

wb.save(out)
print(out)
